from __future__ import annotations

import io
import math
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time

from linecapacity.meter import season_of_month

from ..scheme.editor import EditError
from ..scheme.model import Scheme

ADDRESS_PHRASE = "Адрес прибора АСКУЭ: "
ADDRESS_DIGITS = 16
PERIOD_ROW = 8
FIRST_PERIOD_COLUMN = 8
FIRST_ADDRESS_ROW = 11
ADDRESS_COLUMN = 2
HALF_HOUR = 0.5
NO_LOAD_KW = 1e-15
MAX_BYTES = 5 * 1024 * 1024
WRONG = "Некорректные данные! Повторите ввод!"
INTEGER = re.compile(r"^\s*[+-]?\d+$")
DECIMAL = re.compile(r"^\s*[+-]?(\d+,?\d*|,\d+)([eE][+-]?\d+)?\s*$")


@dataclass
class MeterTable:
    periods: list[str] = field(default_factory=list)
    addresses: list[str] = field(default_factory=list)
    values: list[list[object]] = field(default_factory=list)


@dataclass
class Import:
    scheme: Scheme
    found: int
    season: int | None
    messages: list[str]


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        stamp = value.strftime("%d.%m.%Y")
        return stamp if value.time() == time() else f"{stamp} {value.strftime('%H:%M:%S')}"
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return format(value, ".15g").replace(".", ",")
    return str(value)


def read_table(content: bytes) -> MeterTable:
    if len(content) > MAX_BYTES:
        raise EditError("Файл больше 5 МБ")
    try:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError):
        raise EditError("Ошибка при работе с таблицей Microsoft Excel!") from None
    try:
        sheet = book.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(min_row=1, values_only=True)]
    finally:
        book.close()

    def cell(row: int, column: int) -> object:
        if row - 1 >= len(rows) or column - 1 >= len(rows[row - 1]):
            return None
        return rows[row - 1][column - 1]

    table = MeterTable()
    column = FIRST_PERIOD_COLUMN
    while len(cell_text(cell(PERIOD_ROW, column))) > 2:
        table.periods.append(cell_text(cell(PERIOD_ROW, column)))
        column += 1
    row = FIRST_ADDRESS_ROW
    while cell_text(cell(row, ADDRESS_COLUMN)) != "":
        table.addresses.append(cell_text(cell(row, ADDRESS_COLUMN)))
        table.values.append([cell(row, FIRST_PERIOD_COLUMN + i) for i in range(len(table.periods))])
        row += 1
    if not table.periods:
        raise EditError("В таблице нет получасовых данных: ожидаются даты в строке 8, начиная со столбца H")
    return table


def period_month(text: str) -> int | None:
    start = text.find(".") + 1
    digits = text[start:start + 2]
    return int(digits) if digits.isdigit() else None


class BrokenAddress(ValueError):
    pass


def meter_address(type_text: str) -> int:
    start = len(ADDRESS_PHRASE)
    if len(type_text) <= start:
        return 0
    digits = ""
    for position in range(start, start + ADDRESS_DIGITS):
        if position >= len(type_text):
            raise BrokenAddress(type_text)
        char = type_text[position]
        if char == ",":
            return int(digits) if INTEGER.match(digits) else -1
        digits += char
    return 0


def energy(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        number = float(value)
    else:
        text = cell_text(value)
        if not DECIMAL.match(text):
            return None
        number = float(text.strip().replace(",", "."))
    if not math.isfinite(number):
        return None
    return number / HALF_HOUR


def apply_period(scheme: Scheme, table: MeterTable, period: int, cos_phi: float) -> Import:
    if not 0 <= period < len(table.periods):
        raise EditError("Не выбрана строка в таблице!")
    if not 0 <= cos_phi <= 1:
        raise EditError(WRONG)
    messages = []
    month = period_month(table.periods[period])
    season = season_of_month(month) if month and 1 <= month <= 12 else None
    if season is None:
        messages.append("Некорректная дата в таблице Excel!")
    changed = scheme.model_copy(deep=True)
    found = 0
    address = 0
    wanted = 0
    for number, text in enumerate(table.addresses):
        broken = not INTEGER.match(text)
        if broken:
            messages.append(f"Некорректное значение адреса прибора АСКУЭ в строке №: {number}")
        else:
            wanted = int(text)
        for consumer in changed.consumers:
            if address == -1:
                break
            try:
                address = meter_address(consumer.type_text)
            except BrokenAddress:
                name = consumer.address or consumer.label
                raise EditError(f"Импорт прерван: в дополнительных данных потребителя «{name}» после фразы "
                                f"«{ADDRESS_PHRASE.strip()}» нет запятой, адрес прибора не читается. "
                                "Схема не изменена.") from None
            if address == -1:
                messages.append("Некорректно читается адрес прибора АСКУЭ!")
                break
            if address == wanted:
                found += 1
                consumer.load_type = 1
                power = energy(table.values[number][period])
                if power is None:
                    consumer.p_kw = NO_LOAD_KW
                else:
                    consumer.cos_phi = cos_phi
                    consumer.p_kw = power if power != 0 else NO_LOAD_KW
                break
        if broken:
            break
    messages.append(f"Найдены нагрузки для {found} потребителей из {len(changed.consumers)} потребителей в схеме")
    return Import(changed, found, season, messages)
