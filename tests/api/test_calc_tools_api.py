from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from tests.helpers import ROOT, sign_up, upgrade
from voltplan.app import create_app
from voltplan.exchange import cir

MIXED = ROOT / "tests" / "data" / "calc" / "разные-типы.cir"


@pytest.fixture(scope="module")
def app(tmp_path_factory):
    return create_app(tmp_path_factory.mktemp("site"))


@pytest.fixture
def client(app):
    client = TestClient(app)
    client.card = sign_up(client)
    upgrade(app, client.card["id"])
    return client


def scheme() -> dict:
    return cir.load_bytes(MIXED.read_bytes()).model_dump(mode="json")


def make_admin(app, user_id: int) -> None:
    with app.state.site.store.writing(app.state.site.store.work) as db:
        db.execute("UPDATE profiles SET is_admin = 1 WHERE user_id = ?", (user_id,))


def test_settings_form_lists_every_setting(client):
    data = client.get("/api/calc/settings").json()
    keys = [item["key"] for item in data["fields"]]
    assert len(keys) == len(set(keys)) == 30
    assert set(keys) < set(data["reference"])
    assert data["reference"]["line_air_winter"] == data["defaults"]["line_air_winter"]


def test_calc_takes_settings_and_refuses_wrong_ones(client):
    base = client.post("/api/calc", json={"scheme": scheme()}).json()
    cold = client.post("/api/calc", json={"scheme": scheme(), "settings": {"line_air_winter": -25}}).json()
    assert cold["seasons"]["1"]["max_temperature"] < base["seasons"]["1"]["max_temperature"]
    wrong = client.post("/api/calc", json={"scheme": scheme(), "settings": {"home_winter_cos": 3}})
    assert wrong.status_code == 409 and wrong.json() == {"detail": "Неверные данные! Повторите ввод!"}


def test_reference_is_changed_only_by_administrator(app, client):
    assert client.put("/api/admin/reference", json={"values": {"line_air_winter": -5}}).status_code == 403
    before = client.post("/api/calc", json={"scheme": scheme()}).json()["seasons"]["1"]["max_temperature"]
    make_admin(app, client.card["id"])
    changed = client.put("/api/admin/reference", json={"values": {"line_air_winter": -5}})
    assert changed.status_code == 200 and changed.json()["values"]["line_air_winter"] == -5
    after = client.post("/api/calc", json={"scheme": scheme()}).json()["seasons"]["1"]["max_temperature"]
    assert after < before
    assert client.get("/api/calc/settings").json()["reference"]["line_air_winter"] == -5
    assert client.delete("/api/admin/reference").json()["values"]["line_air_winter"] != -5


def test_project_keeps_settings_and_scenarios(client):
    project = client.post("/api/projects", json={"name": "Сценарии", "scheme": scheme()}).json()
    path = f"/api/projects/{project['id']}"
    assert client.put(f"{path}/settings", json={"settings": {"extra_load_enabled": False}}).status_code == 200
    assert client.get(path).json()["settings"] == {"extra_load_enabled": False}
    listed = client.post(f"{path}/scenarios", json={"name": "Мороз", "settings": {"line_air_winter": -30},
                                                    "period": 1, "min_load": False}).json()
    assert listed == [{"id": listed[0]["id"], "name": "Мороз", "settings": {"line_air_winter": -30.0}, "period": 1,
                       "min_load": False}]
    number = listed[0]["id"]
    renamed = client.patch(f"{path}/scenarios/{number}", json={"name": "Сильный мороз",
                                                                "settings": {"line_air_winter": -35}, "period": 1})
    assert renamed.json()[0]["name"] == "Сильный мороз"
    wrong = client.post(f"{path}/scenarios", json={"name": "Плохой", "settings": {"voltage_min": -1}})
    assert wrong.status_code == 400
    rows = client.post("/api/calc/scenarios", json={
        "scheme": scheme(), "settings": {"extra_load_enabled": False},
        "scenarios": [{"name": "Обычный"}, {"name": "Мороз", "settings": {"line_air_winter": -35}, "period": 1}],
    }).json()["rows"]
    assert [row["name"] for row in rows] == ["Обычный", "Мороз"]
    assert [season["season"] for season in rows[0]["seasons"]] == ["лето", "зима"]
    assert [season["season"] for season in rows[1]["seasons"]] == ["зима"]
    assert rows[1]["seasons"][0]["max_temperature"] < rows[0]["seasons"][1]["max_temperature"]
    assert client.delete(f"{path}/scenarios/{number}").json() == []


def test_load_kind_meter_and_allowed_power(client):
    typical = client.post("/api/calc/load-kind", json={"scheme": scheme(), "typical": True}).json()
    assert typical["message"] == "Типовые значения нагрузки присвоены!"
    spread = client.post("/api/calc/meter", json={"scheme": scheme(), "p_kw": 15, "q_kvar": 4, "month": 8}).json()
    assert spread["season"] == 0 and spread["message"] == "Нагрузки присвоены!"
    assert sum(c["p_kw"] for c in spread["scheme"]["consumers"]) == pytest.approx(15)
    found = client.post("/api/calc/allowed", json={"scheme": scheme(), "index": 0}).json()
    assert found["report"][0] == "Допустимая мощность потребителя по адресу"
    wrong = client.post("/api/calc/meter", json={"scheme": scheme(), "p_kw": 1, "q_kvar": 1, "month": 13})
    assert wrong.status_code == 422


def test_excel_upload_reads_periods_and_applies_one(client):
    book = Workbook()
    sheet = book.active
    sheet.cell(row=8, column=8, value="15.12.2025 17:30")
    sheet.cell(row=8, column=9, value="15.12.2025 18:00")
    sheet.cell(row=11, column=2, value=55)
    sheet.cell(row=11, column=8, value=0.8)
    sheet.cell(row=11, column=9, value=1.1)
    buffer = io.BytesIO()
    book.save(buffer)
    content = buffer.getvalue()
    data = scheme()
    data["consumers"][0]["type_text"] = "Адрес прибора АСКУЭ: 55, мощность по ТУ: 10 кВт"
    periods = client.post("/api/calc/excel/periods", files={"file": ("аскуэ.xlsx", content)}).json()
    assert periods == {"periods": ["15.12.2025 17:30", "15.12.2025 18:00"], "meters": 1}
    applied = client.post("/api/calc/excel/apply", files={"file": ("аскуэ.xlsx", content)},
                          data={"scheme": json.dumps(data), "period": "1", "cos_phi": "0.95"}).json()
    assert applied["found"] == 1 and applied["season"] == 1
    assert applied["scheme"]["consumers"][0]["p_kw"] == pytest.approx(2.2)
    assert applied["settings"] == {}
    warm = client.post("/api/calc/excel/apply", files={"file": ("аскуэ.xlsx", content)},
                       data={"scheme": json.dumps(data), "period": "1", "cos_phi": "0.95", "air": "-12.5",
                             "settings": json.dumps({"home_winter_kw": 2.5})}).json()
    assert warm["settings"] == {"home_winter_kw": 2.5, "line_air_winter": -12.5}
    broken = client.post("/api/calc/excel/periods", files={"file": ("x.xlsx", b"PK-not-really")})
    assert broken.status_code == 409


def test_variants_are_compared_and_the_best_is_chosen(client):
    base = scheme()
    weak = cir.load_bytes(MIXED.read_bytes())
    for line in weak.lines:
        line.length_m *= 6
        line.r_phase_ohm *= 6
        line.r_neutral_ohm *= 6
    project = client.post("/api/projects", json={"name": "Сравнение", "scheme": base}).json()
    main = project["variants"][0]["id"]
    second = client.post(f"/api/projects/{project['id']}/variants", json={"name": "Длинные пролёты"}).json()
    client.put(f"/api/projects/{project['id']}/scheme",
               json={"scheme": weak.model_dump(mode="json"), "variant": second["variant"]})
    result = client.post(f"/api/projects/{project['id']}/compare",
                         json={"variants": [second["variant"], main]}).json()
    names = [row["name"] for row in result["rows"]]
    assert names == ["Длинные пролёты", "Основной"]
    assert result["best"] == 1
    long_row, short_row = result["rows"]
    assert short_row["min_voltage"] > long_row["min_voltage"]
    assert result["best_values"]["min_voltage"] == short_row["min_voltage"]
    assert short_row["quality"]["nominal"] == 220
    assert short_row["annual"]["total_kwh"] > 0
    too_many = client.post(f"/api/projects/{project['id']}/compare", json={"variants": list(range(1, 12))})
    assert too_many.status_code == 409


def test_profile_runs_from_bus_to_consumer(client):
    data = client.post("/api/calc/profile", json={"scheme": scheme(), "consumer": 1, "season": 1}).json()
    points = data["points"]
    assert points[0]["name"] == "Шины ТП" and points[0]["distance_m"] == 0
    assert [p["distance_m"] for p in points] == sorted(p["distance_m"] for p in points)
    consumer = cir.load_bytes(MIXED.read_bytes()).consumers[1]
    assert points[-1]["name"] == consumer.address
    assert data["limits"] == {"low": 198.0, "high": 242.0}


def test_epure_and_losses_need_maximum_tier(app):
    demo = TestClient(app)
    sign_up(demo)
    for path in ("/api/calc/profile", "/api/calc/losses"):
        refused = demo.post(path, json={"scheme": scheme()})
        assert refused.status_code == 403 and "Максимум" in refused.json()["detail"]


@pytest.mark.parametrize("values, message", [
    ({"voltage_rated": 1e160}, "не может быть больше 1 000 000"),
    ({"voltage_min": 230}, "Номинальное напряжение должно быть больше наименьшего"),
    ({"voltage_loss_limit": 150}, "Допустимые потери напряжения"),
    ({"extra_load_kw": -5}, "не могут быть отрицательными"),
    ({"home_minimal_kw": -1}, "не могут быть отрицательными"),
])
def test_unreasonable_settings_are_refused_before_saving(client, values, message):
    project = client.post("/api/projects", json={"name": "Пределы", "scheme": scheme()}).json()
    saved = client.put(f"/api/projects/{project['id']}/settings", json={"settings": values})
    assert saved.status_code == 400 and message in saved.json()["detail"]
    assert client.get(f"/api/projects/{project['id']}/settings").json()["settings"] == {}
    calc = client.post("/api/calc", json={"scheme": scheme(), "settings": values})
    assert calc.status_code == 409 and message in calc.json()["detail"]
    client.delete(f"/api/projects/{project['id']}")


def test_meter_with_not_a_number_is_refused(client):
    for bad in ("NaN", "Infinity"):
        body = '{"scheme": ' + json.dumps(scheme()) + ', "p_kw": ' + bad + ', "q_kvar": 1, "month": 1}'
        answer = client.post("/api/calc/meter", content=body, headers={"Content-Type": "application/json"})
        assert answer.status_code in (409, 422) and "detail" in answer.json()


def test_overflow_in_calculation_is_explained(client, monkeypatch):
    def overflow(*args, **kwargs):
        raise OverflowError("math range error")

    monkeypatch.setattr("voltplan.api.calc.run", overflow)
    answer = client.post("/api/calc", json={"scheme": scheme()})
    assert answer.status_code == 409 and answer.json()["detail"].startswith("Расчёт невозможен")


def test_menu_functions_keep_the_demo_size_limit(app):
    demo = TestClient(app)
    sign_up(demo)
    big = cir.load_bytes((ROOT / "schemes" / "Ботвиново ГКТП397.cir").read_bytes()).model_dump(mode="json")
    for path, body in [("/api/calc/load-kind", {"scheme": big, "typical": True}),
                       ("/api/calc/meter", {"scheme": big, "p_kw": 10, "q_kvar": 1, "month": 1}),
                       ("/api/calc", {"scheme": big})]:
        answer = demo.post(path, json=body)
        assert answer.status_code == 403 and "не больше 40 опор" in answer.json()["detail"], path
    small = demo.post("/api/calc/load-kind", json={"scheme": scheme(), "typical": True})
    assert small.status_code == 200
    book = Workbook()
    book.active.cell(row=8, column=8, value="15.12.2025 17:30")
    book.active.cell(row=11, column=2, value=1)
    buffer = io.BytesIO()
    book.save(buffer)
    excel = demo.post("/api/calc/excel/apply", files={"file": ("аскуэ.xlsx", buffer.getvalue())},
                      data={"scheme": json.dumps(big), "period": "0", "cos_phi": "1"})
    assert excel.status_code == 403 and "не больше 40 опор" in excel.json()["detail"]


def test_typical_loads_return_the_line_temperatures_to_the_reference(client):
    layer = {"line_air_summer": 40.0, "line_air_winter": -30.0, "home_winter_kw": 2.5}
    typical = client.post("/api/calc/load-kind", json={"scheme": scheme(), "typical": True, "settings": layer}).json()
    assert typical["settings"] == {"home_winter_kw": 2.5}
    individual = client.post("/api/calc/load-kind", json={"scheme": scheme(), "typical": False,
                                                          "settings": layer}).json()
    assert individual["settings"] == layer
