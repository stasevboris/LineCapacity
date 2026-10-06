from __future__ import annotations

import io
import math
from datetime import datetime

import pytest
from openpyxl import Workbook

from linecapacity.allowed import allowed_power
from linecapacity.meter import season_of_month, spread_by_meter
from linecapacity.settings import CalculationSettings
from linecapacity.solver import calculate
from tests.helpers import ROOT
from voltplan.calc import settings as calc_settings
from voltplan.calc.bridge import to_model
from voltplan.calc.results import run
from voltplan.calc.tools import allowed, assign_load_kind, spread
from voltplan.exchange import cir
from voltplan.exchange.meters import BrokenAddress, apply_period, meter_address, read_table
from voltplan.scheme.editor import EditError

MIXED = ROOT / "tests" / "data" / "calc" / "разные-типы.cir"
KAMENKA = ROOT / "schemes" / "Каменка КТП158.cir"


def load(path):
    return cir.load_bytes(path.read_bytes())


@pytest.mark.parametrize("month, season", [(1, 1), (4, 1), (5, 0), (9, 0), (10, 1), (12, 1)])
def test_summer_is_may_to_september(month, season):
    assert season_of_month(month) == season


def test_meter_spreads_measured_power_in_proportion_to_typical_loads():
    scheme = load(MIXED)
    settings = CalculationSettings()
    season, loads = spread_by_meter(to_model(scheme), settings, 20.0, 5.0, 1)
    assert season == 1
    peaks = {0: settings.home_winter_kw, 1: settings.electric_home_winter_kw, 2: settings.street_light_winter_kw}
    norms = {0: settings.home_annual_kwh, 1: settings.electric_home_annual_kwh, 2: settings.street_light_annual_kwh}
    others = [c for c in scheme.consumers if c.category == 3]
    expected = sum(peaks[int(c.category)] * c.annual_kwh / norms[int(c.category)] for c in scheme.consumers
                   if c.category in peaks) + sum(c.annual_kwh for c in others) / 8760
    factor = 20.0 / expected
    for consumer, load_ in zip(scheme.consumers, loads, strict=True):
        if consumer.category == 3:
            assert load_.p_kw == pytest.approx(sum(c.annual_kwh for c in others) / (8760 * len(others)) * factor)
        else:
            share = peaks[int(consumer.category)] * factor * consumer.annual_kwh / norms[int(consumer.category)]
            assert load_.p_kw == pytest.approx(share)
        assert load_.cos_phi == pytest.approx(math.cos(math.atan(5 / 20)))
    assert sum(load_.p_kw for load_ in loads) == pytest.approx(20.0)


def test_meter_without_annual_consumption_counts_consumers():
    scheme = load(MIXED)
    _, loads = spread_by_meter(to_model(scheme), CalculationSettings(), 10.0, 0.0, 7, by_annual=False)
    assert sum(load_.p_kw for load_ in loads) == pytest.approx(10.0)
    homes = [load_.p_kw for c, load_ in zip(scheme.consumers, loads, strict=True) if c.category == 0]
    assert len(set(round(value, 9) for value in homes)) == 1
    assert all(load_.cos_phi == 1.0 for load_ in loads)


def test_meter_refuses_zero_power():
    with pytest.raises(EditError, match="Некорректные данные! Повторите ввод!"):
        spread(load(MIXED), CalculationSettings(), 0.0, 1.0, 1, True)


def test_spread_makes_every_consumer_individual():
    scheme, season, message = spread(load(MIXED), CalculationSettings(), 12.0, 3.0, 6, True)
    assert season == 0 and message == "Нагрузки присвоены!"
    assert {c.load_type for c in scheme.consumers} == {1}


@pytest.mark.parametrize("typical", [True, False])
def test_assign_load_kind_changes_only_load_kind(typical):
    scheme = load(MIXED)
    changed, message = assign_load_kind(scheme, typical)
    assert {c.load_type for c in changed.consumers} == {0 if typical else 1}
    assert message == ("Типовые значения нагрузки присвоены!" if typical else
                       "Индивидуальные значения нагрузки присвоены!")
    for before, after in zip(scheme.consumers, changed.consumers, strict=True):
        assert before.model_dump(exclude={"load_type"}) == after.model_dump(exclude={"load_type"})


@pytest.mark.parametrize("text, address", [
    ("Адрес прибора АСКУЭ: 12345, мощность по ТУ: 10 кВт", 12345),
    ("Адрес прибора АСКУЭ: 7,", 7),
    ("без прибора учёта", 0),
    ("Адрес прибора АСКУЭ: 1234567890123456789", 0),
    ("Адрес прибора АСКУЭ:x", 0),
    ("", 0),
    ("Адрес прибора АСКУЭ: 12а3, мощность", -1),
])
def test_meter_address_is_read_from_consumer_data(text, address):
    assert meter_address(text) == address


@pytest.mark.parametrize("text", ["мощность по ТУ: 3,5 кВт", "Адрес прибора АСКУЭ: 8476"])
def test_text_without_comma_after_the_phrase_cannot_be_read(text):
    with pytest.raises(BrokenAddress):
        meter_address(text)
    scheme = load(MIXED)
    scheme.consumers[0].type_text = text
    table = read_table(workbook(["01.07.2026 12:00"], [(8476, [1])]))
    with pytest.raises(EditError, match="Импорт прерван"):
        apply_period(scheme, table, 0, 1.0)


def workbook(periods: list[object], rows: list[tuple[object, list[object]]]) -> bytes:
    book = Workbook()
    sheet = book.active
    for column, period in enumerate(periods, start=8):
        sheet.cell(row=8, column=column, value=period)
    for number, (address, values) in enumerate(rows):
        sheet.cell(row=11 + number, column=2, value=address)
        for column, value in enumerate(values, start=8):
            sheet.cell(row=11 + number, column=column, value=value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def with_addresses(scheme, addresses: list[int | None]):
    changed = scheme.model_copy(deep=True)
    for consumer in changed.consumers:
        consumer.type_text = "без прибора учёта"
    for consumer, address in zip(changed.consumers, addresses, strict=False):
        if address is not None:
            consumer.type_text = f"Адрес прибора АСКУЭ: {address}, мощность по ТУ: 10 кВт"
    return changed


def test_excel_half_hour_energy_becomes_power():
    content = workbook([datetime(2026, 1, 12, 18, 0), datetime(2026, 1, 12, 18, 30)],
                       [(101, [0.4, 0.6]), (102, [1.25, None]), (999, [3, 3])])
    table = read_table(content)
    assert table.periods == ["12.01.2026 18:00:00", "12.01.2026 18:30:00"]
    scheme = with_addresses(load(MIXED), [101, 102, None, None])
    result = apply_period(scheme, table, 1, 0.9)
    assert result.found == 2 and result.season == 1
    first, second = result.scheme.consumers[:2]
    assert (first.load_type, first.p_kw, first.cos_phi) == (1, 1.2, 0.9)
    assert second.p_kw == 1e-15
    assert result.messages == ["Найдены нагрузки для 2 потребителей из 4 потребителей в схеме"]
    assert result.scheme.consumers[2] == scheme.consumers[2]


def test_excel_text_dates_and_values_follow_the_same_rules():
    content = workbook(["01.07.2026 12:00"], [("201", ["0,25"])])
    table = read_table(content)
    result = apply_period(with_addresses(load(MIXED), [201]), table, 0, 1.0)
    assert result.season == 0 and result.scheme.consumers[0].p_kw == 0.5


def test_excel_bad_address_reuses_the_previous_address_then_stops():
    scheme = with_addresses(load(MIXED), [301])
    table = read_table(workbook(["01.07.2026 12:00"], [("х1", [1]), (301, [1])]))
    result = apply_period(scheme, table, 0, 1.0)
    assert result.messages[0] == "Некорректное значение адреса прибора АСКУЭ в строке №: 0"
    assert result.found == 1
    assert result.scheme.consumers[0].p_kw == scheme.consumers[0].p_kw
    assert result.scheme.consumers[1].p_kw == 2.0 and result.scheme.consumers[1].load_type == 1
    table = read_table(workbook(["01.07.2026 12:00"], [(301, [1]), ("12a", [3]), (301, [5])]))
    result = apply_period(scheme, table, 0, 1.0)
    assert result.found == 2 and result.scheme.consumers[0].p_kw == 6.0
    assert result.messages[-1].startswith("Найдены нагрузки для 2 потребителей")


def test_excel_text_with_a_dot_is_not_a_number():
    scheme = with_addresses(load(MIXED), [301, 302, 303])
    table = read_table(workbook(["01.07.2026 12:00"], [(301, ["0.75"]), (302, ["0,75"]), (303, [" 1,5 "])]))
    result = apply_period(scheme, table, 0, 0.9)
    first, second, third = result.scheme.consumers[:3]
    assert first.p_kw == 1e-15 and first.cos_phi == scheme.consumers[0].cos_phi
    assert second.p_kw == 1.5 and second.cos_phi == 0.9
    assert third.p_kw == 3.0
    table = read_table(workbook(["01.07.2026 12:00"], [(301, ["1,5E-1"]), (302, ["2e1"]), (303, ["1.5E-1"])]))
    first, second, third = apply_period(scheme, table, 0, 0.9).scheme.consumers[:3]
    assert first.p_kw == pytest.approx(0.3) and second.p_kw == pytest.approx(40.0) and third.p_kw == 1e-15


def test_excel_refuses_foreign_files_and_bad_cos():
    with pytest.raises(EditError, match="Ошибка при работе с таблицей Microsoft Excel!"):
        read_table(b"not an excel file")
    table = read_table(workbook(["01.07.2026 12:00"], [(1, [1])]))
    with pytest.raises(EditError, match="Некорректные данные! Повторите ввод!"):
        apply_period(load(MIXED), table, 0, 1.5)
    with pytest.raises(EditError, match="Не выбрана строка в таблице!"):
        apply_period(load(MIXED), table, 3, 1.0)


def test_settings_are_checked_like_the_settings_window():
    assert calc_settings.build({}) == CalculationSettings()
    changed = calc_settings.build({"line_air_winter": -15, "extra_load_enabled": False})
    assert changed.line_air_winter == -15 and changed.extra_load_enabled is False
    for wrong in ({"home_summer_cos": 1.2}, {"voltage_min": 0}, {"contact_resistance_ohm": -1},
                  {"home_annual_kwh": 0}, {"street_light_winter_kw": -1}):
        with pytest.raises(EditError, match="Неверные данные! Повторите ввод!"):
            calc_settings.build(wrong)
    with pytest.raises(EditError, match="Неизвестная настройка"):
        calc_settings.build({"surprise": 1})
    with pytest.raises(EditError, match="48"):
        calc_settings.build({"home_summer_curve": [1.0] * 10})


def test_settings_change_the_calculation():
    scheme = load(MIXED)
    base = run(scheme)
    cold = run(scheme, settings=calc_settings.build({"line_air_winter": -20}))
    assert cold["seasons"]["1"]["max_temperature"] < base["seasons"]["1"]["max_temperature"]
    assert cold["seasons"]["0"] == base["seasons"]["0"]
    heavy = run(scheme, settings=calc_settings.build({"home_winter_kw": 3.0}))
    assert heavy["seasons"]["1"]["p_total_kw"] > base["seasons"]["1"]["p_total_kw"]
    without = run(scheme, settings=calc_settings.build({"extra_load_enabled": False}))
    assert without["seasons"]["1"]["extra_load_kw"] == 0


def test_allowed_power_is_found_where_voltage_reaches_the_limit():
    scheme = load(MIXED)
    model = to_model(scheme)
    settings = CalculationSettings()
    base = calculate(model, settings, 2, False)
    index = 0
    found = allowed_power(model, settings, index, base)
    assert found.by_voltage_w > 0 and found.by_transformer_w > 0
    loaded = model.consumers[index]
    consumer = scheme.consumers[index].model_copy()
    consumer.load_type = 1
    consumer.p_kw = found.candidates_w[0] / 1000.0
    changed = scheme.model_copy(deep=True)
    changed.consumers[index] = consumer
    after = calculate(to_model(changed), settings, 2, False).seasons[found.season]
    phase = loaded.phase_no - 1 if loaded.phase_mode == 1 else min(range(3),
                                                                 key=lambda p: base.seasons[found.season]
                                                                 .phase_voltages[index][p])
    floor = abs(after.bus_voltages[phase]) * (1 - settings.voltage_loss_limit / 100)
    assert after.phase_voltages[index][phase] == pytest.approx(floor, abs=2.0)


def test_allowed_power_needs_both_seasons_and_sufficient_capacity():
    with pytest.raises(EditError, match="Потребитель не найден"):
        allowed(load(MIXED), CalculationSettings(), 99)
    bad = load(ROOT / "tests" / "data" / "calc" / "сильная-перегрузка.cir")
    with pytest.raises(EditError, match="пропускная способность сети достаточна"):
        allowed(bad, CalculationSettings(), 0)


def test_allowed_power_rows_follow_the_results_window():
    found = allowed(load(MIXED), CalculationSettings(), 1)
    address = load(MIXED).consumers[1].address
    assert found["report"][:2] == ["Допустимая мощность потребителя по адресу", f"{address}:"]
    assert found["report"][2].startswith(" по критерию допустимых потерь напряжения: ")
    assert found["memo"][:2] == ["", "Допустимая активная мощность потребителя: "]
    assert found["memo"][3].endswith(" кВт")
