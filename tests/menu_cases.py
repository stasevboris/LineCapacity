from __future__ import annotations

import io
import json

from openpyxl import Workbook

from tests.helpers import ROOT

DATA = ROOT / "tests" / "data" / "calc"
RECORDED = DATA / "меню-linecapacity.json"
MIXED = "tests/data/calc/разные-типы.cir"
KAMENKA = "schemes/Каменка КТП158.cir"
GATSKOE = "schemes/Гатское КТП179.cir"
METERED = "tests/data/calc/аскуэ.cir"
BOOK = "tests/data/calc/аскуэ.xlsx"
BOOK_DOT = "tests/data/calc/аскуэ-точка.xlsx"
BOOK_BAD = "tests/data/calc/аскуэ-неверный-адрес.xlsx"
ADDRESSES = {0: 101, 1: 102, 2: 103}
PERIODS = ["15.01.2026 18:00", "15.01.2026 18:30", "15.07.2026 12:00"]
READINGS = [(101, [1.2, 0.8, 0.3]), (102, [None, 2.5, 0.0]), (999, [3.0, 3.0, 3.0]), (103, ["0,75", 1.1, 0.4])]
READINGS_DOT = [(101, ["1.2", 0.8, 0.3]), (102, ["0.75", 2.5, 0.0]), (103, [" 1,5 ", 1.1, 0.4])]
READINGS_BAD = [(101, [1.2, 0.8, 0.3]), ("12a", [2.4, 2.5, 0.0]), (103, [0.75, 1.1, 0.4])]
BOOKS = {BOOK: READINGS, BOOK_DOT: READINGS_DOT, BOOK_BAD: READINGS_BAD}

LOAD_KIND = [{"name": "индивидуальные", "scheme": MIXED, "typical": False},
             {"name": "типовые", "scheme": MIXED, "typical": True}]
METER = [{"name": "зима-с-потреблением", "scheme": KAMENKA, "p_kw": 45.0, "q_kvar": 12.0, "month": 1,
          "by_annual": True},
         {"name": "лето-без-потребления", "scheme": KAMENKA, "p_kw": 30.5, "q_kvar": 4.0, "month": 7,
          "by_annual": False}]
EXCEL = [{"name": "январь-18-00", "scheme": METERED, "book": BOOK, "period": 0, "cos_phi": 0.95, "air": -12.0},
         {"name": "текст-с-точкой", "scheme": METERED, "book": BOOK_DOT, "period": 0, "cos_phi": 0.97, "air": -5.0},
         {"name": "неверный-адрес", "scheme": METERED, "book": BOOK_BAD, "period": 0, "cos_phi": 0.9, "air": 0.0}]
SETTINGS = [{"name": "свои-настройки", "scheme": KAMENKA, "period": 2, "min_load": True,
             "values": {"line_air_winter": -20.0, "transformer_air_winter": -10.0, "home_winter_kw": 2.1,
                        "home_winter_cos": 0.95, "extra_load_kw": 3.5, "electric_home_annual_kwh": 9000.0,
                        "voltage_loss_limit": 6.0, "contact_resistance_ohm": 0.002, "high_voltage_kv": 10.5,
                        "street_light_winter_kw": 0.9}},
            {"name": "без-утяжеления-и-потребления", "scheme": MIXED, "period": 2, "min_load": True,
             "values": {"extra_load_enabled": False, "scale_by_annual": False, "home_summer_kw": 1.2,
                        "line_air_summer": 35.0}}]
ALLOWED = [{"name": "Гатское-0", "scheme": GATSKOE, "consumer": 0},
           {"name": "Гатское-7", "scheme": GATSKOE, "consumer": 7},
           {"name": "разные-типы-1", "scheme": MIXED, "consumer": 1},
           {"name": "разные-типы-3", "scheme": MIXED, "consumer": 3}]


def book_bytes(readings: list | None = None) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.cell(row=1, column=1, value="Выгрузка АСКУЭ")
    for column, period in enumerate(PERIODS, start=8):
        sheet.cell(row=8, column=column, value=period)
    for number, (address, values) in enumerate(READINGS if readings is None else readings):
        sheet.cell(row=11 + number, column=2, value=address)
        for column, value in enumerate(values, start=8):
            sheet.cell(row=11 + number, column=column, value=value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


KINDS = ("load_kind", "meter", "excel", "settings", "allowed")


def recorded() -> dict:
    stored = json.loads(RECORDED.read_text(encoding="utf-8")) if RECORDED.exists() else {}
    return {kind: stored.get(kind, []) for kind in KINDS}


def number_or_text(text: str):
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return text


def same_file(ours: bytes, theirs: bytes, tolerance: float = 1e-13) -> list[str]:
    first = ours.decode("cp1251").splitlines()
    second = theirs.decode("cp1251").splitlines()
    if len(first) != len(second):
        return [f"строк: VoltPlan {len(first)}, LineCapacity {len(second)}"]
    problems = []
    for number, (mine, their) in enumerate(zip(first, second, strict=True), 1):
        if mine == their:
            continue
        a, b = number_or_text(mine), number_or_text(their)
        if isinstance(a, float) and isinstance(b, float) and abs(a - b) <= tolerance * max(abs(a), abs(b)):
            continue
        problems.append(f"строка {number}: VoltPlan «{mine}» — LineCapacity «{their}»")
    return problems
