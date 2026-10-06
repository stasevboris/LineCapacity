from __future__ import annotations

import io
import re
import time
from urllib.parse import urlparse

import pytest
from openpyxl import Workbook

from tests.helpers import PASSWORD, ROOT, fresh_email
from tests.live.ui import EditorPage, Server
from voltplan.exchange import cir

pytestmark = pytest.mark.live

SHOTS = ROOT / "logs" / "screens"
KAMENKA = ROOT / "schemes" / "Каменка КТП158.cir"
MIXED = ROOT / "tests" / "data" / "calc" / "разные-типы.cir"
FUTURE = time.strftime("%m/%y", time.localtime(time.time() + 400 * 86400))
CYRILLIC = re.compile(r"[А-Яа-яЁё]")


@pytest.fixture(scope="module")
def server():
    with Server() as running:
        yield running


@pytest.fixture(scope="module")
def browser():
    playwright_api = pytest.importorskip("playwright.sync_api")
    with playwright_api.sync_playwright() as playwright:
        try:
            instance = playwright.chromium.launch(channel="chrome")
        except Exception:
            try:
                instance = playwright.chromium.launch()
            except Exception as error:
                pytest.skip(f"браузер недоступен: {error}")
        yield instance
        instance.close()


def shot(page, name: str) -> None:
    SHOTS.mkdir(parents=True, exist_ok=True)
    page.screenshot(path=str(SHOTS / f"{name}.png"))


def watch(page) -> list[str]:
    errors: list[str] = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.on("console", lambda message: errors.append(message.text) if message.type == "error"
            and not re.search(r"status of (400|401|402|403|404|409|429)", message.text) else None)
    page.on("dialog", lambda dialog: dialog.accept())
    return errors


def person(server, browser, tier: str = "max", name: str = "Проверка", width: int = 1440):
    context = browser.new_context(viewport={"width": width, "height": 900}, accept_downloads=True)
    email = fresh_email("site")
    card = server.sign_up(context, email, name, tier)
    page = context.new_page()
    errors = watch(page)
    return context, page, errors, {**card, "email": email}


def visible_text(page, selector: str = "main") -> str:
    return page.evaluate("""(selector) => {
        const copy = document.querySelector(selector).cloneNode(true);
        copy.querySelectorAll('[translate="no"], #profile-language').forEach((node) => node.remove());
        return copy.innerText;
    }""", selector)


def without_user_data(text: str, path) -> str:
    scheme = cir.load_bytes(path.read_bytes())
    own = [scheme.transformer.label, scheme.transformer.type_name]
    for group in (scheme.lines, scheme.poles, scheme.consumers):
        for item in group:
            own += [getattr(item, name, "") for name in ("label", "address", "type_text", "type_name")]
    for value in sorted((value for value in own if value), key=len, reverse=True):
        text = text.replace(value, "")
    return text


def arrive(page, path: str) -> None:
    page.wait_for_url(lambda url: urlparse(url).path == path)
    page.wait_for_load_state("load")


def open_file(page, path) -> None:
    lines = len(cir.load_bytes(path.read_bytes()).lines)
    page.set_input_files("#file-cir", str(path))
    page.wait_for_function(f"document.querySelectorAll('#layer-lines [data-kind=\"line\"]').length === {lines}")


def calculate(page, period: int = 2) -> None:
    page.keyboard.press("F7")
    page.wait_for_selector("#modal:not([hidden])")
    page.locator(f'#modal-body input[name="period"][value="{period}"]').check()
    with page.expect_response(lambda response: response.url.endswith("/api/calc")):
        page.click("#modal-ok")
    page.wait_for_timeout(300)
    if page.locator("#report").is_visible():
        page.click("#report-close")


def calc_menu(page, item: str) -> None:
    page.click("#btn-calc-menu")
    page.click(f"#{item}")


def test_registration_login_and_protected_pages(server, browser):
    context = browser.new_context(viewport={"width": 1280, "height": 860})
    page = context.new_page()
    errors = watch(page)
    page.goto(server.url + "app")
    assert "/login?next=/app" in page.url
    page.goto(server.url + "register?next=/account")
    email = fresh_email("reg")
    page.fill("#name", "Ирина Мельник")
    page.fill("#email", email)
    page.fill("#password", "коротко")
    page.fill("#password2", "коротко")
    page.click("#btn-submit")
    page.wait_for_function("document.getElementById('form-error').textContent.length > 0")
    assert page.locator("#form-error").inner_text() == "Пароль должен быть не короче 8 знаков"
    page.fill("#password", PASSWORD)
    page.fill("#password2", PASSWORD)
    page.click("#btn-submit")
    arrive(page, "/account")
    page.wait_for_function("document.getElementById('tier-name').textContent === 'Демо'")
    assert page.locator("#who-name").inner_text() == "Ирина Мельник"
    assert page.locator("#acc-email").inner_text() == email.lower()
    shot(page, "30-кабинет-после-регистрации")
    assert page.locator(".nav-toggle").is_visible() and not page.locator("#logout").is_visible()
    page.click(".nav-toggle")
    box = page.locator("#logout").bounding_box()
    assert page.evaluate("([x, y]) => document.elementFromPoint(x, y).id", [box["x"] + 5, box["y"] + 5]) == "logout"
    page.click("#logout")
    page.wait_for_url(server.url)
    page.goto(server.url + "login?next=/app")
    page.fill("#email", email.upper())
    page.fill("#password", PASSWORD)
    page.click("#btn-submit")
    arrive(page, "/app")
    assert page.locator("#empty").is_visible()
    assert errors == []
    context.close()


def test_project_is_saved_and_opened_again(server, browser):
    context, page, errors, _ = person(server, browser)
    page.goto(server.url + "app")
    editor = EditorPage(page)
    open_file(page, KAMENKA)
    page.keyboard.press("Control+s")
    page.wait_for_selector("#modal:not([hidden])")
    page.fill('#modal-body input', "Каменка, КТП 158")
    page.click("#modal-ok")
    page.wait_for_selector("#project-bar:not([hidden])")
    assert page.locator("#project-title").inner_text() == "Каменка, КТП 158"
    project = int(re.search(r"project=(\d+)", page.url).group(1))
    editor.update("consumer", 0, {"label": "д.1а"}, zoom=6)
    assert page.locator("#dirty").is_visible()
    page.click("#btn-save-project")
    page.wait_for_selector("#dirty", state="hidden")
    page.goto(server.url + "account#projects")
    page.wait_for_selector(f'#projects-list [data-project="{project}"]')
    page.click(f'[data-project="{project}"] a.btn-primary')
    arrive(page, "/app")
    page.wait_for_function("document.querySelectorAll('#layer-consumers [data-kind]').length > 5")
    stored = page.evaluate("(id) => fetch(`/api/projects/${id}/scheme`).then((r) => r.json())", project)
    assert stored["scheme"]["consumers"][0]["label"] == "д.1а"
    shot(page, "31-проект-открыт")
    assert editor.errors == [] and errors == []
    context.close()


def test_team_reader_sees_project_but_cannot_change_it(server, browser):
    owner_context, owner, owner_errors, _ = person(server, browser, "max", "Владелец")
    reader_context, reader, reader_errors, reader_card = person(server, browser, "demo", "Читатель")
    owner.goto(server.url + "app")
    open_file(owner, MIXED)
    owner.keyboard.press("Control+s")
    owner.wait_for_selector("#modal:not([hidden])")
    owner.fill('#modal-body input', "Общий проект")
    owner.click("#modal-ok")
    owner.wait_for_selector("#project-bar:not([hidden])")
    project = int(re.search(r"project=(\d+)", owner.url).group(1))
    owner.goto(server.url + "collab#teams")
    owner.fill("#team-name", "Бригада")
    owner.click("#team-new button")
    owner.wait_for_selector(".team")
    owner.wait_for_selector("#invite-target option:nth-child(2)", state="attached")
    owner.fill("#invite-email", reader_card["email"])
    owner.select_option("#invite-target", index=1)
    owner.select_option("#invite-role", "reader")
    owner.click("#invite-form button")
    owner.wait_for_selector("#outgoing .status-waiting")
    reader.goto(server.url + "collab#invitations")
    reader.wait_for_function("document.getElementById('bell-count').textContent === '1'")
    reader.click("#incoming button.btn-accept")
    reader.wait_for_selector("#incoming .status-accepted")
    shot(reader, "31а-совместная-работа")
    owner.goto(server.url + "account#projects")
    owner.click(f'#projects-list [data-project="{project}"] button:has-text("Команде")')
    owner.click('#dialog-body button:has-text("Открыть доступ")')
    owner.wait_for_timeout(400)
    reader.goto(server.url + f"app?project={project}")
    reader.wait_for_selector("#readonly-badge:not([hidden])")
    assert reader.locator("#btn-save-project").is_disabled()
    assert reader.locator("#actions .action").count() == 0
    calculate(reader)
    assert reader.locator("#results-bar").is_visible()
    shot(reader, "32-проект-читателя")
    status = reader.evaluate("""(id) => fetch('/api/projects/' + id + '/scheme').then((r) => r.json())
        .then((d) => fetch('/api/projects/' + id + '/scheme', {method: 'PUT', headers: {'Content-Type':
        'application/json'}, body: JSON.stringify({scheme: d.scheme})})).then((r) => r.status)""", project)
    assert status == 403
    assert owner_errors == [] and reader_errors == []
    owner_context.close()
    reader_context.close()


def test_friends_exchange_messages_with_a_file(server, browser):
    first_context, first, first_errors, first_card = person(server, browser, "demo", "Анна")
    second_context, second, second_errors, second_card = person(server, browser, "demo", "Борис")
    first.goto(server.url + "collab#friends")
    first.fill("#friend-email", second_card["email"])
    first.click("#friend-invite button")
    first.wait_for_selector("#outgoing .status-waiting")
    second.goto(server.url + "collab#invitations")
    second.click("#incoming button.btn-accept")
    second.wait_for_selector("#friends-list [data-friend]")
    first.goto(server.url + "collab#chat")
    first.click(f'[data-chat="user-{second_card["id"]}"]')
    first.fill("#chat-text", "Посмотри схему, пожалуйста")
    first.set_input_files("#chat-file", {"name": "Каменка.cir", "mimeType": "application/octet-stream",
                                         "buffer": KAMENKA.read_bytes()})
    assert first.locator("#chat-attach-name").inner_text() == "Каменка.cir"
    first.click("#chat-form button[type=submit]")
    first.wait_for_selector("#chat-feed .bubble.mine")
    second.goto(server.url + "collab#chat")
    second.click(f'[data-chat="user-{first_card["id"]}"]')
    second.wait_for_selector("#chat-feed .bubble a.file")
    with second.expect_download() as download:
        second.click("#chat-feed .bubble a.file")
    assert download.value.path().read_bytes() == KAMENKA.read_bytes()
    shot(second, "33-переписка-с-файлом")
    assert first_errors == [] and second_errors == []
    first_context.close()
    second_context.close()


def test_calc_menu_settings_scenarios_and_allowed_power(server, browser):
    context, page, errors, _ = person(server, browser)
    page.goto(server.url + "app")
    open_file(page, MIXED)
    calculate(page)
    before = page.evaluate("document.querySelectorAll('#layer-results text').length")
    calc_menu(page, "tool-settings")
    page.wait_for_selector("#modal:not([hidden]) .setting-row")
    page.fill('.setting-row[data-key="line_air_winter"] input', "-15")
    assert page.locator('.setting-row[data-key="line_air_winter"]').get_attribute("class").endswith("changed")
    shot(page, "34-настройки-расчёта")
    page.click("#modal-ok")
    page.wait_for_selector("#modal", state="hidden")
    calculate(page)
    assert before > 0
    page.click('#layer-consumers [data-kind="consumer"][data-index="0"]', force=True)
    page.wait_for_selector("#btn-allowed:not([disabled])")
    with page.expect_response(lambda response: "/api/calc/allowed" in response.url):
        page.click("#btn-allowed")
    page.wait_for_selector("#report:not([hidden])")
    rows = page.locator("#report-body .report-row").all_inner_texts()
    assert rows[0] == "Допустимая мощность потребителя по адресу"
    assert "Допустимая активная мощность потребителя:" in page.locator("#props .memo").inner_text()
    shot(page, "35-допустимая-мощность")
    page.click("#report-close")
    calc_menu(page, "tool-meter")
    page.select_option('#modal-body select[name="month"]', "1")
    page.fill('#modal-body input[name="p_kw"]', "12")
    page.fill('#modal-body input[name="q_kvar"]', "3")
    page.click("#modal-ok")
    page.wait_for_selector("#modal", state="hidden")
    assert page.locator("#toast").inner_text() == "Нагрузки присвоены!"
    page.keyboard.press("F7")
    page.wait_for_selector("#modal:not([hidden])")
    assert page.locator('#modal-body input[name="period"][value="1"]').is_checked()
    page.click("#modal-cancel")
    assert errors == []
    context.close()


def excel_file() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.cell(row=8, column=8, value="10.01.2026 18:00")
    sheet.cell(row=11, column=2, value=4401)
    sheet.cell(row=11, column=8, value=1.5)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_excel_import_and_scenarios_in_project(server, browser):
    context, page, errors, _ = person(server, browser)
    scheme = cir.load_bytes(MIXED.read_bytes())
    scheme.consumers[0].type_text = "Адрес прибора АСКУЭ: 4401, мощность по ТУ: 10 кВт"
    path = SHOTS.parent / "аскуэ-схема.cir"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(cir.dump_bytes(scheme))
    page.goto(server.url + "app")
    open_file(page, path)
    calc_menu(page, "tool-excel")
    page.set_input_files('#modal-body input[type="file"]', {"name": "аскуэ.xlsx", "buffer": excel_file(),
                                                            "mimeType": "application/octet-stream"})
    page.wait_for_selector('#modal-body select[name="period"] option')
    page.click("#modal-ok")
    page.wait_for_selector("#modal", state="hidden")
    assert page.locator("#toast").inner_text().startswith("Найдены нагрузки для 1 потребителей")
    page.keyboard.press("Control+s")
    page.wait_for_selector("#modal:not([hidden])")
    page.fill('#modal-body input', "Сценарии")
    page.click("#modal-ok")
    page.wait_for_selector("#project-bar:not([hidden])")
    project = int(re.search(r"project=(\d+)", page.url).group(1))
    for name, value in (("Обычная зима", None), ("Сильный мороз", "-30")):
        page.evaluate("""([id, name, value]) => fetch('/api/projects/' + id + '/scenarios', {method: 'POST',
            headers: {'Content-Type': 'application/json'}, body: JSON.stringify({name, period: 1,
            settings: value ? {line_air_winter: Number(value)} : {}})})""", [project, name, value])
    calc_menu(page, "tool-scenarios")
    page.wait_for_selector('#modal-body [data-scenario]')
    page.click('.scenario-tools button.btn-primary')
    page.wait_for_selector(".summary-table")
    names = page.locator(".summary-table tbody td:first-child").all_inner_texts()
    assert [name for name in names if name] == ["Обычная зима", "Сильный мороз"]
    shot(page, "36-сценарии")
    page.click("#modal-ok")
    assert errors == []
    context.close()


def test_variants_are_compared_in_a_table(server, browser):
    context, page, errors, _ = person(server, browser)
    page.goto(server.url + "app")
    open_file(page, MIXED)
    page.keyboard.press("Control+s")
    page.wait_for_selector("#modal:not([hidden])")
    page.fill('#modal-body input', "Варианты")
    page.click("#modal-ok")
    page.wait_for_selector("#project-bar:not([hidden])")
    page.click("#btn-variant-new")
    page.wait_for_selector("#modal:not([hidden])")
    page.fill('#modal-body input', "Без утяжеления")
    page.click("#modal-ok")
    page.wait_for_function("document.getElementById('variant-select').options.length === 2")
    calc_menu(page, "tool-compare")
    page.wait_for_selector("#modal:not([hidden]) input[type=checkbox]")
    page.click("#modal-ok")
    page.wait_for_selector(".compare-table")
    assert page.locator(".compare-table tr.winner").count() == 1
    assert page.locator(".compare-table td.best").count() >= 1
    shot(page, "37-сравнение-вариантов")
    page.click("#modal-ok")
    assert errors == []
    context.close()


def test_epure_losses_and_docx_follow_the_tier(server, browser):
    context, page, errors, _ = person(server, browser, "max")
    page.goto(server.url + "app")
    open_file(page, KAMENKA)
    calc_menu(page, "tool-profile")
    page.wait_for_selector("#modal:not([hidden])")
    page.click("#modal-ok")
    page.wait_for_selector("svg.epure path.epure-line")
    assert page.locator("svg.epure path.epure-line").count() >= 1
    shot(page, "38-эпюра")
    page.click("#modal-ok")
    calc_menu(page, "tool-losses")
    page.wait_for_selector(".losses-table")
    assert "Потери энергии за год" in page.locator(".losses-table").inner_text()
    page.click("#modal-ok")
    with page.expect_download() as download:
        calc_menu(page, "tool-docx")
    assert download.value.path().read_bytes()[:2] == b"PK"
    demo_context, demo, demo_errors, _ = person(server, browser, "demo")
    demo.goto(server.url + "app")
    open_file(demo, MIXED)
    calc_menu(demo, "tool-profile")
    demo.click("#modal-ok")
    demo.wait_for_function("document.getElementById('modal-error').textContent.includes('Максимум')")
    assert errors == [] and demo_errors == []
    context.close()
    demo_context.close()


def test_demo_card_payment_upgrades_the_tier(server, browser):
    context, page, errors, _ = person(server, browser, "demo")
    page.goto(server.url + "tariffs")
    page.wait_for_selector('.tier-card[data-tier="pro"] a.btn-primary')
    shot(page, "39-тарифы")
    shot(page, "39-тарифы")
    page.click('.tier-card[data-tier="pro"] a.btn-primary')
    arrive(page, "/pay")
    page.fill("#card-number", "4000000000000002")
    page.fill("#card-expiry", FUTURE.replace("/", ""))
    page.fill("#card-cvv", "123")
    page.click("#btn-pay")
    page.wait_for_selector(".outcome.bad")
    page.goto(server.url + "pay?tier=pro")
    page.fill("#card-number", "4111111111111111")
    assert page.locator("#card-number").input_value() == "4111 1111 1111 1111"
    assert page.locator("#card-brand").inner_text() == "Visa"
    page.fill("#card-expiry", FUTURE.replace("/", ""))
    page.fill("#card-cvv", "123")
    page.fill("#card-holder", "ivan petrov")
    page.click("#btn-pay")
    page.wait_for_selector(".outcome.good")
    shot(page, "40-оплата-прошла")
    page.goto(server.url + "account#payments")
    page.wait_for_selector("#payments-list table")
    page.wait_for_function("document.getElementById('tier-name').textContent === 'Профессионал'")
    statuses = page.locator("#payments-list tbody tr").count()
    assert statuses == 2
    assert errors == []
    context.close()


def test_administrator_changes_tier_and_reference(server, browser):
    context, page, errors, card = person(server, browser, "demo", "Администратор")
    server.app.state.site.store.work.execute("UPDATE profiles SET is_admin = 1 WHERE user_id = ?", (card["id"],))
    other_context, other, other_errors, other_card = person(server, browser, "demo", "Инженер")
    page.goto(server.url + "admin#users")
    page.fill("#users-query", other_card["email"])
    page.click("#users-search button")
    page.wait_for_selector(f'[data-user="{other_card["id"]}"]')
    row = page.locator(f'[data-user="{other_card["id"]}"]')
    row.locator("select").select_option("max")
    row.locator('input[type="date"]').fill(time.strftime("%Y-%m-%d", time.localtime(time.time() + 20 * 86400)))
    row.locator("button").click()
    page.wait_for_timeout(400)
    other.goto(server.url + "account#payments")
    other.wait_for_function("document.getElementById('tier-name').textContent === 'Максимум'")
    assert other.locator("#term").is_visible()
    page.goto(server.url + "admin#overview")
    page.wait_for_function("document.getElementById('s-users').textContent !== '—'")
    assert int(page.locator("#s-max").inner_text()) >= 1
    shot(page, "41а-администрирование-обзор")
    page.goto(server.url + "admin#reference")
    page.wait_for_selector('#reference-fields [name="line_air_winter"]')
    page.fill('#reference-fields [name="line_air_winter"]', "-7")
    page.click('#reference-form button[type="submit"]')
    page.wait_for_timeout(400)
    reference = page.evaluate("fetch('/api/calc/settings').then((r) => r.json()).then((d) => d.reference)")
    assert reference["line_air_winter"] == -7
    shot(page, "41-администрирование")
    page.evaluate("fetch('/api/admin/reference', {method: 'DELETE'})")
    assert errors == [] and other_errors == []
    context.close()
    other_context.close()


@pytest.mark.parametrize("code, word", [("en", "Open editor"), ("zh", "打开编辑器")])
def test_interface_speaks_three_languages(server, browser, code, word):
    context, page, errors, _ = person(server, browser, "max")
    page.goto(server.url)
    page.click(f'#language button[data-lang="{code}"]')
    page.wait_for_load_state("load")
    page.wait_for_selector(f'#language button[data-lang="{code}"].on')
    page.wait_for_timeout(500)
    text = page.locator("main").inner_text()
    assert not CYRILLIC.search(text), CYRILLIC.findall(text)[:20]
    assert word in text
    shot(page, f"42-главная-{code}")
    page.goto(server.url + "app")
    open_file(page, MIXED)
    calculate(page)
    page.click("#btn-calc-menu")
    menu = page.locator("#calc-menu").inner_text()
    assert not CYRILLIC.search(menu)
    page.keyboard.press("Escape")
    page.click("#btn-report")
    page.wait_for_selector("#report:not([hidden])")
    report = page.locator("#report-body").inner_text()
    assert "winter" in report.lower() or "冬" in report
    shot(page, f"43-редактор-{code}")
    page.click("#report-close")
    editor = EditorPage(page)
    for kind in ("consumer", "line", "transformer"):
        editor.click_object(kind, 0)
        page.wait_for_selector("#props .memo-line")
        shown = without_user_data(page.locator("#props").inner_text(), MIXED)
        assert not CYRILLIC.search(shown), (kind, shown[:400])
    shot(page, f"43а-свойства-{code}")
    project = page.evaluate("""(scheme) => fetch('/api/projects', {method: 'POST', headers: {'Content-Type':
        'application/json'}, body: JSON.stringify({name: 'Проект', scheme})}).then((r) => r.json())""",
                            cir.load_bytes(MIXED.read_bytes()).model_dump(mode="json"))
    page.evaluate("""(id) => fetch('/api/projects/' + id + '/variants', {method: 'POST', headers: {'Content-Type':
        'application/json'}, body: JSON.stringify({name: 'Второй'})})""", project["id"])
    page.goto(server.url + f"app?project={project['id']}")
    page.wait_for_function("document.getElementById('variant-select').options.length === 2")
    main = page.evaluate("document.getElementById('variant-select').options[0].text")
    assert main == {"en": "Main", "zh": "主方案"}[code], main
    page.click("#btn-variant-rename")
    assert page.input_value("#modal input") == main
    page.click("#modal-ok")
    page.wait_for_selector("#modal", state="hidden")
    assert page.evaluate("document.getElementById('variant-select').options[0].text") == main
    calc_menu(page, "tool-compare")
    page.wait_for_selector("#modal:not([hidden]) input[type=checkbox]")
    assert not CYRILLIC.search(visible_text(page, "#modal")), visible_text(page, "#modal")[:300]
    page.click("#modal-ok")
    page.wait_for_selector(".compare-table")
    assert not CYRILLIC.search(visible_text(page, "#modal")), visible_text(page, "#modal")[:400]
    shot(page, f"43б-сравнение-{code}")
    page.click("#modal-ok")
    page.goto(server.url + "account#projects")
    page.wait_for_selector("#projects-list table, #projects-list .empty-note")
    assert not CYRILLIC.search(visible_text(page)), CYRILLIC.findall(visible_text(page))[:20]
    page.goto(server.url + "collab")
    page.wait_for_selector("#teams-list .empty-note")
    assert not CYRILLIC.search(visible_text(page)), CYRILLIC.findall(visible_text(page))[:20]
    page.click('#language button[data-lang="ru"]')
    page.wait_for_load_state("load")
    assert errors == []
    context.close()


def test_narrow_phone_screen_has_no_sideways_scroll(server, browser):
    context = browser.new_context(viewport={"width": 390, "height": 844})
    page = context.new_page()
    errors = watch(page)
    for path in ("", "features", "about", "faq", "status", "tariffs", "login", "register"):
        page.goto(server.url + path)
        page.wait_for_timeout(200)
        assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth"), path
    shot(page, "44-телефон-регистрация")
    page.click(".nav-toggle")
    shown = page.evaluate("""() => [...document.querySelectorAll('.nav .links > *')]
        .filter((el) => getComputedStyle(el).display !== 'none' && el.getBoundingClientRect().height > 0)
        .map((el) => el.getAttribute('data-auth') || '-')""")
    assert "in" not in shown and shown.count("out") == 2, shown
    box = page.locator(".nav .links a[data-auth='out']").last.bounding_box()
    assert page.evaluate("([x, y]) => Boolean(document.elementFromPoint(x, y).closest('.nav'))",
                         [box["x"] + 5, box["y"] + 5])
    assert errors == []
    context.close()
